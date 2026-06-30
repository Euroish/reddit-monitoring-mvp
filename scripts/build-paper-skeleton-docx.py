from __future__ import annotations

import json
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUT_DIR = ROOT / "output" / "paper-draft-2026-05-27"
OUT_PATH = OUT_DIR / "Reddit中AliExpress相关讨论的消费者风险感知与购买决策支持研究_论文骨架.docx"
RAW_META = ROOT / "output" / "reddit-research-raw-dataset-2026-05-26" / "raw-dataset-metadata.json"
MECH_SUMMARY = ROOT / "output" / "reddit-research-matched-2026-05-26" / "deeper-risk-mechanism-summary.json"


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, v in {"top": top, "start": start, "bottom": bottom, "end": end}.items():
        node = tc_mar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")


def set_table_width(table, widths):
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    for row in table.rows:
        for idx, width in enumerate(widths):
            cell = row.cells[idx]
            cell.width = Inches(width)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            set_cell_margins(cell)


def add_hyperlink(paragraph, text: str, url: str):
    part = paragraph.part
    r_id = part.relate_to(
        url,
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True,
    )
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), r_id)
    run = OxmlElement("w:r")
    r_pr = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "0563C1")
    r_pr.append(color)
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    r_pr.append(underline)
    run.append(r_pr)
    text_node = OxmlElement("w:t")
    text_node.text = text
    run.append(text_node)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


def style_run(run, bold=False, italic=False, size=None, color=None):
    run.bold = bold
    run.italic = italic
    if size:
        run.font.size = Pt(size)
    if color:
        run.font.color.rgb = RGBColor.from_string(color)
    run.font.name = "SimSun"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "SimSun")


def add_para(doc, text="", style=None, bold_prefix=None):
    p = doc.add_paragraph(style=style)
    if bold_prefix and text.startswith(bold_prefix):
        r1 = p.add_run(bold_prefix)
        style_run(r1, bold=True)
        r2 = p.add_run(text[len(bold_prefix):])
        style_run(r2)
    else:
        r = p.add_run(text)
        style_run(r)
    return p


def add_bullets(doc, items):
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        r = p.add_run(item)
        style_run(r)


def add_numbered(doc, items):
    for item in items:
        p = doc.add_paragraph(style="List Number")
        r = p.add_run(item)
        style_run(r)


def add_heading(doc, text: str, level: int):
    p = doc.add_heading(level=level)
    r = p.add_run(text)
    style_run(r, bold=True)
    return p


def configure_styles(doc: Document):
    section = doc.sections[0]
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "SimSun"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "SimSun")
    normal.font.size = Pt(11)
    normal.paragraph_format.line_spacing = 1.333
    normal.paragraph_format.space_after = Pt(8)

    for name, size, color, before, after in [
        ("Heading 1", 16, "2E74B5", 18, 10),
        ("Heading 2", 13, "2E74B5", 12, 6),
        ("Heading 3", 12, "1F4D78", 8, 4),
    ]:
        style = styles[name]
        style.font.name = "SimSun"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "SimSun")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)


def add_source_note(doc, label: str, url: str):
    p = doc.add_paragraph()
    r = p.add_run(label + " ")
    style_run(r)
    add_hyperlink(p, url, url)


def main() -> None:
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    raw_meta = json.loads(RAW_META.read_text(encoding="utf-8"))
    mech = json.loads(MECH_SUMMARY.read_text(encoding="utf-8"))

    doc = Document()
    configure_styles(doc)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = title.add_run("Reddit中AliExpress相关讨论的消费者风险感知与购买决策支持研究")
    style_run(r, bold=True, size=18, color="0B2545")
    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = subtitle.add_run("——基于帖子-评论匹配数据的主题分析")
    style_run(r, size=14, color="1F4D78")

    meta = doc.add_paragraph()
    meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = meta.add_run("论文骨架与初稿框架 | 数据日期：2026-05-26 | 版本：基础结构稿")
    style_run(r, size=10, color="555555")

    add_heading(doc, "选题判断", 1)
    add_para(
        doc,
        "建议沿用当前题目，不建议转向 AliExpress 与 Temu/SHEIN 的比较研究，也不建议写广告效果或因果购买行为。当前数据最强的证据链是：AliExpress 相关帖子提出风险、困惑或经验，评论区围绕这些帖子提供解释、建议、警示和风险降低策略。因此，论文应定位为 Reddit 社区语境下的探索性文本研究，而不是总体消费者行为或因果效应研究。",
    )
    add_para(
        doc,
        "推荐题目：Reddit中AliExpress相关讨论的消费者风险感知与购买决策支持研究——基于帖子-评论匹配数据的主题分析",
        bold_prefix="推荐题目：",
    )

    add_heading(doc, "摘要（待完善）", 1)
    add_para(
        doc,
        "本文以 Reddit 中 AliExpress 相关帖子及评论为研究对象，基于帖子-评论匹配数据，探讨跨境电商语境下消费者风险感知的主要类型及评论区在购买决策支持中的作用。研究拟结合主题分析、关键词辅助编码与情感倾向分析，识别物流时间、税费政策、退款争议、商品质量、卖家与平台信任、价格价值权衡等核心议题，并进一步分析评论区如何通过经验分享、风险提示、支付保护建议、卖家筛选建议和购买时机判断参与消费者决策支持。研究发现预计将表明，AliExpress 相关讨论并非单纯负面评价，而是在低价价值与多重风险之间形成持续权衡；Reddit 评论区则承担了电子口碑与风险降低的信息功能。",
    )
    add_para(doc, "关键词：AliExpress；Reddit；消费者风险感知；购买决策支持；电子口碑；主题分析")

    add_heading(doc, "一、引言", 1)
    add_heading(doc, "1.1 研究背景", 2)
    add_para(
        doc,
        "跨境电商平台降低了消费者接触海外商品的门槛，但也放大了物流时间、关税税费、商品质量、售后退款和卖家可信度等不确定性。AliExpress 作为典型跨境电商平台，既具有低价、品类丰富、可替代性强等优势，也经常被消费者置于风险权衡之中。",
    )
    add_para(
        doc,
        "Reddit 作为以兴趣社区为核心的公共讨论平台，聚集了大量围绕平台购物经验、商品评价、售后争议和购买建议的用户生成内容。这些内容适合观察消费者如何表达风险、如何寻求同伴意见，以及评论区如何提供购买决策支持。",
    )
    add_heading(doc, "1.2 研究问题", 2)
    add_numbered(
        doc,
        [
            "Reddit 中 AliExpress 相关讨论呈现出哪些主要消费者风险感知类型？",
            "评论区如何通过经验分享、警示、建议和替代方案参与消费者购买决策支持？",
            "低价价值判断如何与物流、税费、质量、退款和信任风险共同影响用户讨论？",
        ],
    )
    add_heading(doc, "1.3 研究意义", 2)
    add_bullets(
        doc,
        [
            "理论层面：将感知风险理论与电子口碑研究结合，解释 Reddit 评论区在跨境购物风险降低中的作用。",
            "方法层面：使用帖子-评论匹配数据，而不是单独分析帖子或评论，提高文本解释的上下文完整性。",
            "实践层面：为跨境电商平台、卖家和品牌理解消费者风险关注点、信任建构和购买阻力提供参考。",
        ],
    )

    add_heading(doc, "二、文献综述", 1)
    add_heading(doc, "2.1 消费者感知风险与跨境电商购物", 2)
    add_para(
        doc,
        "感知风险理论认为，消费者在购买决策中会评估潜在损失和不确定性。跨境电商语境下，风险并不限于商品本身，还包括物流时间、税费政策、售后维权、平台规则、卖家可信度和支付安全。本研究可将风险类型初步划分为政策税费风险、物流时间风险、服务补救风险、商品质量风险、价格价值不确定性和平台信任风险。",
    )
    add_heading(doc, "2.2 电子口碑与在线社区中的购买决策支持", 2)
    add_para(
        doc,
        "电子口碑研究关注消费者生成内容如何影响信息采纳、信任形成和消费决策。Reddit 的帖子和评论具有明显的同伴经验分享特征。Bonifazi 等人关于 Reddit eWOM power 的研究说明 Reddit 帖子可以作为电子口碑传播和影响力分析对象；Noguti 对 Reddit 内容社区的研究则表明，帖子语言特征与用户互动存在关系。这些研究为使用 Reddit 帖子和评论分析消费者讨论提供了方法依据。",
    )
    add_heading(doc, "2.3 Reddit 作为消费讨论与平台知识空间", 2)
    add_para(
        doc,
        "Reddit 社区不仅承载品牌或平台评价，也形成围绕消费经验、规则解释、风险提示和操作建议的知识流。Kwon 与 Shao 对 Reddit 电商相关社区的研究表明，Reddit 可作为平台知识和社区治理问题的观察场域。本文不研究非法电商或治理问题，但借鉴其将 Reddit 视为平台知识流动空间的思路。",
    )
    add_heading(doc, "2.4 社交媒体文本情感与主题分析方法", 2)
    add_para(
        doc,
        "Hutto 与 Gilbert 提出的 VADER 模型面向社交媒体文本情感分析，适合短文本、口语化和带有强调语气的在线评论。本文后续可将 VADER 作为基础情感测量工具，并结合人工复核或 LLM 辅助编码，以降低俚语、反讽和多语种文本带来的误差。",
    )

    add_heading(doc, "三、数据与方法", 1)
    add_heading(doc, "3.1 数据来源", 2)
    add_para(
        doc,
        "本文使用 Reddit 中 AliExpress 相关帖子和评论数据。原始数据集已完成去重和帖子-评论关联，但尚未完成主题清洗、语言筛选、噪声剔除和人工编码。",
    )
    table = doc.add_table(rows=1, cols=2)
    table.style = "Table Grid"
    headers = ["项目", "当前数据规模"]
    for i, h in enumerate(headers):
        table.rows[0].cells[i].text = h
        set_cell_shading(table.rows[0].cells[i], "F4F6F9")
    data_rows = [
        ("去重帖子数", str(raw_meta["uniquePostCount"])),
        ("去重评论数", str(raw_meta["uniqueCommentCount"])),
        ("有评论匹配的帖子数", str(raw_meta["matchedPostCount"])),
        ("帖子-评论断链数", str(raw_meta["missingLinkCount"])),
        ("建议下一步", "人工清洗 false positives、deleted/removed、无关 megathread、低信息评论和非研究语言样本"),
    ]
    for left, right in data_rows:
        cells = table.add_row().cells
        cells[0].text = left
        cells[1].text = right
    set_table_width(table, [2.0, 4.5])

    add_heading(doc, "3.2 数据清洗计划", 2)
    add_numbered(
        doc,
        [
            "保留 post 与 comment 的对应关系，以 postExternalId/externalId 作为主键。",
            "删除 deleted、removed、空文本、极短无意义评论和明显 bot/meta 评论。",
            "剔除 megathread、giveaway、game link、纯优惠码、明显无关召回等帖子。",
            "标记语言类型；非英语文本可单独保留、翻译或从主分析中排除。",
            "保留帖子标题、正文、subreddit、创建时间、score、numComments、评论正文和评论 score 等字段。",
        ],
    )
    add_heading(doc, "3.3 分析方法", 2)
    add_bullets(
        doc,
        [
            "主题分析：先用关键词与开放编码形成初始主题，再通过人工复核合并主题。",
            "情感分析：可使用 VADER 对英文评论进行基础情感打分，再抽样复核。",
            "帖子-评论联动：以帖子风险类型为起点，分析评论区出现的建议、警示和风险降低策略。",
            "描述性统计：统计主题频次、评论量、score、时间分布和 subreddit 分布。",
        ],
    )

    add_heading(doc, "四、预期分析框架", 1)
    add_para(
        doc,
        "初步分析显示，本研究不应只停留在“风险主题有哪些”，而应进一步解释评论区如何将风险讨论转化为决策支持。可使用如下机制框架：风险触发 -> 评论区解释与经验分享 -> 风险降低策略 -> 购买决策支持。",
    )
    table2 = doc.add_table(rows=1, cols=3)
    table2.style = "Table Grid"
    for i, h in enumerate(["风险类型", "评论区可能出现的决策支持", "可观察指标"]):
        table2.rows[0].cells[i].text = h
        set_cell_shading(table2.rows[0].cells[i], "F4F6F9")
    framework_rows = [
        ("政策/税费风险", "购买时机建议、税费解释、地区经验比较", "tariff、tax、VAT、EU、customs、wait/before/after"),
        ("物流/时间风险", "追踪解释、等待建议、正常化安抚", "shipping、tracking、delivery、delay、package"),
        ("退款/服务失败", "开 dispute、PayPal/chargeback、保留证据", "refund、dispute、PayPal、evidence、proof"),
        ("质量/真伪风险", "卖家筛选、评价查看、同类商品经验", "fake、quality、review、seller、legit"),
        ("价格/价值权衡", "低价风险接受、替代平台比较、是否值得买", "cheap、worth、price、deal、Amazon/Temu/eBay"),
        ("卖家/平台信任", "警示、推荐、平台规则解释", "seller、trust、beware、warning、Ali says"),
    ]
    for row in framework_rows:
        cells = table2.add_row().cells
        for i, value in enumerate(row):
            cells[i].text = value
    set_table_width(table2, [1.55, 2.75, 2.2])

    add_heading(doc, "五、结果章节写作骨架", 1)
    add_heading(doc, "5.1 样本描述", 2)
    add_para(doc, "报告清洗后保留的帖子数、评论数、subreddit 分布、时间分布、帖子互动指标，并说明数据采集范围与局限。")
    add_heading(doc, "5.2 主要风险主题", 2)
    add_para(doc, "按政策税费、物流时间、服务补救、商品质量、价格价值、平台信任六类展开，每类结合帖子和评论说明。")
    add_heading(doc, "5.3 评论区的购买决策支持机制", 2)
    add_para(doc, "分析评论区如何提供等待/购买时机、支付保护、证据保留、卖家筛选、替代平台、风险接受和警示等策略。")
    add_heading(doc, "5.4 风险与价值的共同出现", 2)
    add_para(doc, "重点讨论 AliExpress 相关讨论中的张力：消费者并非只表达负面情绪，而是在低价、可获得性、商品丰富度与风险之间进行权衡。")

    add_heading(doc, "六、讨论", 1)
    add_bullets(
        doc,
        [
            "理论讨论：将感知风险与电子口碑结合，解释 Reddit 评论区作为风险降低信息空间的作用。",
            "实践启示：平台和卖家应减少信息不确定性，强化物流透明、售后规则、卖家信用和支付保护提示。",
            "方法反思：帖子-评论匹配比单独评论分析更适合解释购买决策支持。",
        ],
    )

    add_heading(doc, "七、局限与后续研究", 1)
    add_bullets(
        doc,
        [
            "Reddit 用户不能代表全部 AliExpress 消费者。",
            "搜索式样本可能包含误召回，必须经过人工清洗。",
            "评论抓取有深度和数量限制，不代表完整评论树。",
            "本文不能证明真实购买行为或因果影响，只能解释讨论中的风险感知和决策支持。",
            "后续可加入多平台比较、时间序列或人工编码一致性检验。",
        ],
    )

    add_heading(doc, "八、结论（待撰写）", 1)
    add_para(
        doc,
        "本部分后续应回扣三个研究问题，概括 AliExpress 相关 Reddit 讨论中的主要风险类型、评论区决策支持机制，以及风险与价格价值判断的共同作用。结论必须保持边界：研究对象是 Reddit 语境下的公开讨论，而非全部 AliExpress 消费者。",
    )

    add_heading(doc, "参考文献（初步）", 1)
    refs = [
        ("Bonifazi 等：Modeling, Evaluating, and Applying the eWoM Power of Reddit Posts", "https://www.mdpi.com/2504-2289/7/1/47"),
        ("Noguti：Post language and user engagement in online content communities", "https://doi.org/10.1108/EJM-12-2014-0785"),
        ("Hutto 与 Gilbert：VADER: A Parsimonious Rule-Based Model for Sentiment Analysis of Social Media Text", "https://doi.org/10.1609/icwsm.v8i1.14550"),
        ("Kwon 与 Shao：Dark Knowledge and Platform Governance: A Case of an Illicit E-Commerce Community in Reddit", "https://doi.org/10.1177/0002764221989770"),
    ]
    for label, url in refs:
        add_source_note(doc, label, url)

    add_heading(doc, "附录 A：后续清洗与编码清单", 1)
    add_bullets(
        doc,
        [
            "确认所有分析记录都来自 raw-posts-with-comments.json 或由其派生。",
            "建立清洗字段：keep/exclude、exclude_reason、language、topic_primary、topic_secondary、sentiment、advice_type。",
            "抽样 30-50 个帖子进行人工开放编码，形成主题词典后再批量编码。",
            "每类主题保留少量代表性 paraphrase，避免在论文中暴露用户身份或大段原文。",
        ],
    )

    doc.save(OUT_PATH)
    print(OUT_PATH)


if __name__ == "__main__":
    main()
